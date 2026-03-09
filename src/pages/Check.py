import streamlit as st
from io import BytesIO
from docx import Document
from datetime import datetime
from services.pdf_reader_services import PdfReaderManager
from services.pipeline import PipelineManager
from transformers import pipeline


def load_css():
    try:
        with open("assets/css/styles.css") as f:
            st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)
    except FileNotFoundError:
        pass

load_css()

@st.cache_resource
def load_model():
    return pipeline("text-classification", model="Juner/AI-generated-text-detection")


clf = load_model()

def generate_docx(content, result_label, score):

    doc = Document()

    doc.add_heading('AI Writing Indicator - Analysis Report', 0)

    doc.add_paragraph(
        f"Generated time: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    )

    doc.add_heading('Kết quả phân tích', level=1)

    doc.add_paragraph(f"Label: {result_label}")
    doc.add_paragraph(f"Confidence: {score*100:.2f}%")

    doc.add_heading('Nội dung văn bản', level=1)
    doc.add_paragraph(content)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer

st.set_page_config(page_title="AI Analyzer | OUEL", layout="wide")

st.markdown("""
<div class="header-split">
    <h1 class="title-split">AI <span class="blue-grad">Writing</span> Indicator</h1>
    <div class="status-badge">Engine v2.0 Active</div>
</div>
""", unsafe_allow_html=True)

col_left, col_right = st.columns(2, gap="medium")

with col_left:

    st.markdown('<p class="label-v3">INPUT SOURCE</p>', unsafe_allow_html=True)

    data = st.file_uploader("", type="pdf", label_visibility="collapsed")

    if data:
        data_text = PdfReaderManager.load_file(data)
        write_permission = True
    else:
        data_text = ""
        write_permission = False

    text = st.text_area(
        label="Input Area",
        value=data_text,
        height=500,
        placeholder="Dán văn bản hoặc tải file PDF...",
        disabled=write_permission,
        label_visibility="collapsed"
    )

    check_btn = st.button("PHÂN TÍCH NGAY", use_container_width=True)


with col_right:

    st.markdown('<p class="label-v3">ANALYSIS REPORT</p>', unsafe_allow_html=True)

    result_placeholder = st.container(border=True)

    with result_placeholder:

        if check_btn:

            processed_text = PipelineManager.input_processing(text)

            if not processed_text.strip():
                st.info("Vui lòng nhập nội dung ở bên trái.")

            else:

                with st.spinner("Đang tính toán..."):

                    try:

                        PipelineManager.output_processing(
                            processed_text, clf, st
                        )

                        result = clf(processed_text)

                        if isinstance(result, list):
                            result = result[0]

                        label = result.get("label", "Unknown")
                        score = result.get("score", 0)

                        label_map = {
                            "LABEL_0": "Human Written",
                            "LABEL_1": "AI Generated"
                        }

                        label = label_map.get(label, label)

                        result_label = label

                        st.markdown(
                            '<div style="margin-top:20px;"></div>',
                            unsafe_allow_html=True
                        )

                        docx_file = generate_docx(
                            processed_text,
                            result_label,
                            score
                        )

                        st.download_button(
                            label="XUẤT BÁO CÁO (.DOCX)",
                            data=docx_file,
                            file_name="Kết quả kiểm tra.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True
                        )

                    except Exception as e:
                        st.error(f"Lỗi kết nối mô hình: {e}")

        else:

            st.markdown(
                '<div class="empty-state"><p>Kết quả sẽ hiển thị tại đây.</p></div>',
                unsafe_allow_html=True
            )

st.markdown(
    '<div class="footer-v3"><p>© 2026 OUEL RESEARCH | OPEN UNIVERSITY</p></div>',
    unsafe_allow_html=True
)